# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx"]
# ///
"""deck.html → .pptx, best effort: slide backgrounds, text boxes (paragraphs, size, bold / italic, colour, alignment),
pictures, tables, the common shapes with their fill and outline, straight lines, speaker notes, slide transitions
(PowerPoint's nearest: fade, push, wipe, cover, pull, circle, zoom; Morph as PowerPoint's Morph, shapes named "!!<id>"
so they pair) and builds (appear, fade, wipe in; fade out; on click, with or after the previous one, with their
duration and delay; other effects become a fade). Videos and web embeds are left out.

    uv run glideshow/pptx_export.py DECK/deck.html OUT.pptx

The server runs this for File → Download → PowerPoint; python-pptx stays out of the server itself.
"""
import re, sys
from html.parser import HTMLParser
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

EMU_PX = 914400 * 10 / 1920   # the 1920 px canvas is 10 in wide
PX_PT = 192 / 72
SHAPES = {"rect": "RECTANGLE", "roundrect": "ROUNDED_RECTANGLE", "ellipse": "OVAL", "triangle": "ISOSCELES_TRIANGLE", "diamond": "DIAMOND",
          "parallelogram": "PARALLELOGRAM", "trapezoid": "TRAPEZOID", "pentagon": "REGULAR_PENTAGON", "hexagon": "HEXAGON", "octagon": "OCTAGON",
          "star": "STAR_5_POINT", "cloud": "CLOUD", "arrow": "RIGHT_ARROW", "leftarrow": "LEFT_ARROW", "uparrow": "UP_ARROW", "downarrow": "DOWN_ARROW",
          "leftrightarrow": "LEFT_RIGHT_ARROW", "chevron": "CHEVRON", "callout": "RECTANGULAR_CALLOUT", "roundcallout": "ROUNDED_RECTANGULAR_CALLOUT",
          "cross": "CROSS", "heart": "HEART", "donut": "DONUT", "can": "CAN", "cube": "CUBE", "homeplate": "PENTAGON", "rttriangle": "RIGHT_TRIANGLE"}
BLOCK = {"div", "p", "li", "h1", "h2", "h3", "pre", "tr"}
# every other shape: its OOXML preset name from shapes.js' PRESET_OF table, which PowerPoint knows natively
_js = (next(d for d in (Path(__file__).resolve().parent / "editor", Path(__file__).resolve().parent.parent / "editor") if d.is_dir()) / "deck" / "shapes.js").read_text()
PRESET_OF = dict(re.findall(r"['\"]?([\w-]+)['\"]?:\s*'(\w+)'", _js[_js.index("PRESET_OF = {"):_js.index("};", _js.index("PRESET_OF = {"))]))


def mso(shape_name):
    if shape_name in SHAPES: return getattr(MSO_SHAPE, SHAPES[shape_name])
    try: return MSO_SHAPE.from_xml(PRESET_OF.get(shape_name, ""))
    except (ValueError, KeyError): return None


class Node:
    def __init__(self, tag, attrs, parent=None):
        self.tag, self.attrs, self.parent, self.kids = tag, dict(attrs), parent, []

    def cls(self):
        return self.attrs.get("class", "").split()

    def find_all(self, pred):
        for k in self.kids:
            if isinstance(k, Node):
                if pred(k): yield k
                yield from k.find_all(pred)

    def text(self):
        return "".join(k if isinstance(k, str) else k.text() for k in self.kids)


class Tree(HTMLParser):
    VOID = {"img", "br", "meta", "link", "input", "hr", "source"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = self.cur = Node("root", {})

    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs, self.cur); self.cur.kids.append(n)
        if tag not in self.VOID: self.cur = n

    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag: n = n.parent
        if n is not self.root: self.cur = n.parent

    def handle_data(self, data):
        self.cur.kids.append(data)


def style(n):
    return {k.strip(): v.strip() for k, v in (d.split(":", 1) for d in n.attrs.get("style", "").split(";") if ":" in d)}


def px(v, default=0.0):
    m = re.match(r"(-?[\d.]+)px", v or "")
    return float(m.group(1)) if m else default


def rgb(v):
    m = re.fullmatch(r"#([0-9a-fA-F]{6})(?:[0-9a-fA-F]{2})?", (v or "").strip()) or re.fullmatch(r"#([0-9a-fA-F]{3})", (v or "").strip())
    if not m: return None
    h = m.group(1)
    return RGBColor.from_string(h if len(h) == 6 else "".join(c * 2 for c in h))


def paragraphs(content):
    """[(text, block style)] — one per block child / list item / <br>, inline text joined."""
    out, line = [], []
    def walk(n, inherited):
        for k in n.kids:
            if isinstance(k, str): line.append(k)
            elif k.tag == "br": flush(inherited)
            elif k.tag in ("ul", "ol"): walk(k, inherited)
            elif k.tag in BLOCK: flush(inherited); walk(k, {**inherited, **style(k)}); flush({**inherited, **style(k)}, k.tag == "li")
            elif k.tag not in ("svg", "aside", "script", "style"): walk(k, inherited)
    def flush(st, bullet=False):
        t = re.sub(r"\s+", " ", "".join(line)).strip()
        if t: out.append((("• " if bullet else "") + t, st))
        line.clear()
    walk(content, {}); flush({})
    return out


def add_text(shape, content, box_style):
    st = style(content)
    tf = shape.text_frame; tf.word_wrap = True
    paras = paragraphs(content)
    if not paras: return
    tf.text = ""
    for i, (t, pst) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = t
        size = px(pst.get("font-size") or st.get("font-size"), 37.33)
        r.font.size = Pt(round(size / PX_PT, 1))
        r.font.bold = (pst.get("font-weight") or st.get("font-weight")) in ("bold", "700")
        r.font.italic = (pst.get("font-style") or st.get("font-style")) == "italic"
        c = rgb(pst.get("color") or st.get("color"))
        if c: r.font.color.rgb = c
        fam = (st.get("font-family") or "").split(",")[0].strip("'\" ")
        if fam: r.font.name = fam
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}.get(pst.get("text-align") or st.get("text-align") or "", PP_ALIGN.LEFT if not box_style.get("shape") else PP_ALIGN.CENTER)


def export(deck: Path, out: Path):
    t = Tree(); t.feed(deck.read_text()); root = t.root
    slides_el = next(root.find_all(lambda n: n.tag == "div" and "slides" in n.cls()))
    m = re.match(r"(\d+)x(\d+)$", slides_el.attrs.get("data-size", ""))
    W, H = (int(m.group(1)), int(m.group(2))) if m else (1920, 1080)
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(int(W * EMU_PX)), Emu(int(H * EMU_PX))
    blank = prs.slide_layouts[6]
    sections = [s for top in slides_el.kids if isinstance(top, Node) and top.tag == "section"
                for s in ([k for k in top.kids if isinstance(k, Node) and k.tag == "section"] or [top])]
    for sec in sections:
        slide = prs.slides.add_slide(blank)
        bg = rgb(sec.attrs.get("data-background-color"))
        if bg: slide.background.fill.solid(); slide.background.fill.fore_color.rgb = bg
        built = []
        for box in [k for k in sec.kids if isinstance(k, Node) and "box" in k.cls()]:
            shapes = add_box(slide, box, deck.parent, 0, 0) or []
            for sh in shapes: sh.name = "!!" + (box.attrs.get("data-morph") or box.attrs.get("id") or sh.name)   # PowerPoint's Morph pairs shapes by "!!" names
            if "fragment" in box.cls() and shapes: built.append((box, [sh.shape_id for sh in shapes]))
        transition(slide, sec)
        if built: timing(slide, built)
        notes = [k for k in sec.kids if isinstance(k, Node) and k.tag == "aside" and "notes" in k.cls()]
        if notes: slide.notes_slide.notes_text_frame.text = notes[0].text().strip()
    prs.save(out)


NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
# data-transition → PowerPoint's element (Wipe / Move in / Slide from …: the side the new slide comes from, PowerPoint's dir the way it moves)
TRANSITIONS = {"fade": "<p:fade/>", "fade-in fade-out": "<p:fade/>", "fade-black": '<p:fade thruBlk="1"/>', "slide": '<p:push dir="l"/>', "sl-left": '<p:push dir="r"/>',
               "sl-up": '<p:push dir="u"/>', "sl-down": '<p:push dir="d"/>', "wipe-r": '<p:wipe dir="r"/>', "wipe-l": '<p:wipe dir="l"/>', "wipe-d": '<p:wipe dir="d"/>',
               "wipe-u": '<p:wipe dir="u"/>', "move-in-r": '<p:cover dir="r"/>', "move-in-l": '<p:cover dir="l"/>', "move-in-d": '<p:cover dir="d"/>',
               "move-in-u": '<p:cover dir="u"/>', "reveal-l": '<p:pull dir="l"/>', "iris": "<p:circle/>", "zoom": "<p:zoom/>", "k-scale": "<p:zoom/>", "fade-move": '<p:push dir="l"/>'}


def transition(slide, sec):
    a = sec.attrs; spd = {"slow": "slow", "fast": "fast"}.get(a.get("data-transition-speed", ""), "med")
    if "data-auto-animate" in a:   # Morph: PowerPoint's own (2016 and later), a fade for older ones
        xml = (f'<mc:AlternateContent {NS} xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main"><mc:Choice Requires="p159">'
               f'<p:transition spd="slow" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" p14:dur="{int(float(a.get("data-auto-animate-duration") or 1) * 1000)}"><p159:morph option="byObject"/></p:transition>'
               f'</mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>')
    else:
        kind = a.get("data-transition", "")
        if not kind or kind == "none": return
        xml = f'<p:transition {NS} spd="{spd}">{TRANSITIONS.get(kind, "<p:fade/>")}</p:transition>'
    el = slide._element; el.insert(list(el).index(el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}clrMapOvr")) + 1, parse_xml(xml))


OUTS = {"fade-out", "line-erase", "disappear", "blur-out", "scale-out", "iris-out"}
WIPES = {"wipe-r": ("left", 8), "wipe-l": ("right", 2), "wipe-d": ("up", 1), "wipe-u": ("down", 4)}   # Wipe from left = PowerPoint's wipe(left), subtype


def timing(slide, built):
    """The slide's builds as PowerPoint's main sequence: each click a group; "with previous" joins it, "after previous" waits
    for what came before in it."""
    n = [2]
    def nid(): n[0] += 1; return n[0]
    order = sorted(built, key=lambda b: (int(b[0].attrs.get("data-fragment-index", 0) or 0), "data-with-previous" in b[0].attrs))
    clicks, last = [], None
    for box, ids in order:
        idx, start = int(box.attrs.get("data-fragment-index", 0) or 0), box.attrs.get("data-start", "")
        if clicks and (idx == last or start in ("after", "transition")): clicks[-1].append((box, ids))
        else: clicks.append([(box, ids)])
        last = idx
    groups = []
    for click in clicks:
        at, inner, first = 0, [], True
        for box, ids in click:
            a, eff = box.attrs, next((c for c in box.cls() if c not in ("box", "fragment", "image", "table", "group", "visible")), "fade-in")
            dur, delay = int(float(a.get("data-duration") or 0.5) * 1000), int(float(a.get("data-delay") or 0) * 1000)
            start = a.get("data-start", "")
            if not first and start == "after": at = inner[-1][1]
            node = "clickEffect" if first and start not in ("after", "transition") else ("afterEffect" if start in ("after", "transition") else "withEffect")
            out = eff in OUTS
            cls_, preset, sub, filt = ("exit", 10, 0, "fade") if out else ("entr", 1, 0, None) if eff == "appear" else ("entr", 22, WIPES[eff][1], f"wipe({WIPES[eff][0]})") if eff in WIPES else ("entr", 10, 0, "fade")
            behaviours = ""
            for spid in ids:
                vis = f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{dur - 1 if out else 0}"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="{"hidden" if out else "visible"}"/></p:to></p:set>'
                anim = f'<p:animEffect transition="{"out" if out else "in"}" filter="{filt}"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>' if filt else ""
                behaviours += (anim + vis) if out else (vis + anim)
            effect = f'<p:par><p:cTn id="{nid()}" presetID="{preset}" presetClass="{cls_}" presetSubtype="{sub}" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>{behaviours}</p:childTnLst></p:cTn></p:par>'
            if first or start == "after": inner.append([[effect], at + delay + dur, at])
            else: inner[-1][0].append(effect); inner[-1][1] = max(inner[-1][1], inner[-1][2] + delay + dur)
            first = False
        pars = "".join(f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="{t0}"/></p:stCondLst><p:childTnLst>{"".join(effs)}</p:childTnLst></p:cTn></p:par>' for effs, _, t0 in inner)
        groups.append(f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>{pars}</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(groups)}</p:childTnLst></p:cTn>'
           '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = slide._element; after = [c for c in el if c.tag.split("}")[1] in ("cSld", "clrMapOvr", "transition", "AlternateContent")]
    el.insert(list(el).index(after[-1]) + 1, parse_xml(xml))


def add_box(slide, box, base, ox, oy):
    """Adds the box's PowerPoint shapes; returns them (several for a group), or nothing when it's left out."""
    st, a = style(box), box.attrs
    x, y, w, h = ox + px(st.get("left")), oy + px(st.get("top")), px(st.get("width"), 1), px(st.get("height"), 1)
    E = lambda v: Emu(int(v * EMU_PX))
    content = next((k for k in box.kids if isinstance(k, Node) and "content" in k.cls()), None)
    shape_name, cls = a.get("data-shape", ""), box.cls()
    if "group" in cls:
        return [sh for k in box.kids if isinstance(k, Node) and "box" in k.cls() for sh in (add_box(slide, k, base, x, y) or [])]
    if "image" in cls:
        img = next((n for n in box.find_all(lambda n: n.tag == "img")), None)
        src = img and img.attrs.get("src", "")
        if src and not src.startswith(("data:", "http")) and (base / src).exists():
            return [slide.shapes.add_picture(str(base / src), E(x), E(y), E(w), E(h))]
        return
    if "video" in cls or "embed" in cls:
        return   # PowerPoint can't take these from HTML; left out
    if "table" in cls:
        rows = [r for r in box.find_all(lambda n: n.tag == "tr")]
        cols = max((len([c for c in r.kids if isinstance(c, Node) and c.tag in ("td", "th")]) for r in rows), default=0)
        if rows and cols:
            gf = slide.shapes.add_table(len(rows), cols, E(x), E(y), E(w), E(h)); tb = gf.table
            size = px(style(content).get("font-size") if content else None, 37.33)
            for i, r in enumerate(rows):
                for j, c in enumerate([c for c in r.kids if isinstance(c, Node) and c.tag in ("td", "th")]):
                    cell = tb.cell(i, j); cell.text = re.sub(r"\s+", " ", c.text()).strip()
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs: run.font.size = Pt(round(size / PX_PT, 1))
            return [gf]
        return
    sw = px(st.get("--sw"), 4)
    if shape_name in ("line", "arrowline", "elbow", "curve"):
        up = a.get("data-dir") == "up"
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x), E(y + h if up else y), E(x + w), E(y if up else y + h))
        ln.line.width = Emu(int(sw * EMU_PX)); c = rgb(st.get("--stroke"))
        if c: ln.line.color.rgb = c
        return [ln]
    if shape_name and mso(shape_name) is not None:
        sh = slide.shapes.add_shape(mso(shape_name), E(x), E(y), E(w), E(h))
        f, s = rgb(st.get("--fill")), rgb(st.get("--stroke"))
        if f: sh.fill.solid(); sh.fill.fore_color.rgb = f
        else: sh.fill.background()
        if s: sh.line.color.rgb = s; sh.line.width = Emu(int(sw * EMU_PX))
        else: sh.line.fill.background()
    elif shape_name:
        return   # a shape PowerPoint export doesn't know yet
    else:
        sh = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        f = rgb(st.get("--fill"))
        if f: sh.fill.solid(); sh.fill.fore_color.rgb = f
    if content is not None: add_text(sh, content, {"shape": shape_name})
    return [sh]


if __name__ == "__main__":
    export(Path(sys.argv[1]), Path(sys.argv[2]))
