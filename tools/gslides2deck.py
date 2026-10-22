# /// script
# requires-python = ">=3.10"
# dependencies = ["playwright", "python-pptx"]
# ///
"""Google Slides → deck, with the exact geometry of the original, without a .pptx/.pdf download.

The deck's embed viewer (/presentation/d/ID/embed) ships every slide as an SVG (`SK_svgData`): shapes as paths with
their real fill, outline, dash and coordinates, text as glyph outlines plus the plain text in `aria-label`, pictures
as <image>. This tool turns that into boxes:

  1. Capture (needs the signed-in browser). Print the snippet and run it in a tab on docs.google.com, e.g. with the
     Chrome tool's javascript runner. It fetches the embed page, decodes the slides, inlines the pictures it can
     fetch, and downloads gslides-<id>.json:
         uv run tools/gslides2deck.py snippet
  2. Convert:
         uv run tools/gslides2deck.py convert gslides-<id>.json DECK [--slides 1,4-7] [--keep]
     writes DECK/deck.html (creating it if needed) and DECK/assets/. Existing slides are replaced unless --keep,
     which appends.
  3. Optional, exact references for comparing (the original slides rendered from the capture, 1920 px wide):
         uv run tools/gslides2deck.py refs gslides-<id>.json DECK/refs-svg

Mapping: rectangles, rounded rectangles and ellipses become those shapes; straight lines become line boxes (their
arrowheads exact, grouped with the line); outlines without a fill (curves, polylines, open polygons) become
`data-shape="path"`; any other filled outline becomes the editor preset it matches (block arrows, callouts, stars …, with
their adj values and flips, when the outlines overlap ≥ 96 %) or else a box holding its exact SVG. Text becomes real text: font
family, size, weight, style and line breaks are fitted to the glyph outlines in headless Chrome, then placed so the
baselines land where Slides put them; fits under 0.9 glyph overlap are listed in DECK/source/convert-log.txt to check by eye.
What the SVG doesn't carry: link targets, speaker notes, animations. Pass the deck's own .pptx download (File → Download →
PowerPoint) as --pptx to get the notes and the on-click builds (as reveal fragments, matched to boxes by position).
Animated GIFs stay GIFs; `refs` and test/snap.py show their first frame so comparisons are stable.

"""
import argparse, base64, os, hashlib, html, json, math, re, sys, urllib.request
from pathlib import Path
from xml.etree import ElementTree as ET

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))

SVG_NS = "{http://www.w3.org/2000/svg}"
XLINK = "{http://www.w3.org/1999/xlink}href"

# Families tried when fitting text. Web-safe ones need no link; the rest come from Google Fonts.
WEBSAFE = ["Arial", "Georgia", "Times New Roman", "Courier New", "Verdana", "Trebuchet MS"]
GOOGLE = ["Roboto", "Open Sans", "Lato", "Montserrat", "Merriweather", "Lora", "Playfair Display", "Oswald", "Raleway",
          "Roboto Slab", "Source Sans 3", "Nunito", "Inter", "Poppins", "PT Sans", "EB Garamond", "Comfortaa"]

SNIPPET = r"""
await (async () => {
  const id = location.pathname.split('/')[3];
  const page = await (await fetch(`/presentation/d/${id}/embed?rm=minimal`, { credentials: 'include' })).text();
  const dec = s => s.replace(/\\x([0-9a-f]{2})/gi, (_, x) => String.fromCharCode(parseInt(x, 16)))
    .replace(/\\u([0-9a-f]{4})/gi, (_, x) => String.fromCharCode(parseInt(x, 16))).replace(/\\\//g, '/').replace(/\\\\/g, '\\');
  const slides = [...page.matchAll(/SK_svgData = '((?:[^'\\]|\\.)*)'/g)].map(m => dec(m[1]));
  const urls = [...new Set(slides.flatMap(s => [...s.matchAll(/xlink:href="([^"]+)"/g)].map(m => m[1].replace(/&amp;/g, '&'))))];
  const images = {};
  for (const u of urls) {
    try {
      const r = await fetch(u, { credentials: 'include' });
      if (!r.ok) continue;
      const b = await r.blob();
      images[u] = await new Promise(res => { const f = new FileReader(); f.onload = () => res(f.result); f.readAsDataURL(b); });
    } catch (e) { /* cross-origin: the converter downloads it */ }
  }
  const title = (document.title || '').replace(/ - Google Slides$/, '');
  const blob = new Blob([JSON.stringify({ id, title, source: 'embed', slides, images })], { type: 'application/json' });
  const a = Object.assign(document.createElement('a'), { href: URL.createObjectURL(blob), download: `gslides-${id.slice(0, 8)}.json` });
  document.body.append(a); a.click(); a.remove();
  return { slides: slides.length, images: Object.keys(images).length, of: urls.length, bytes: blob.size, file: a.download };
})()
"""


# ---------------------------------------------------------------- SVG paths

def parse_path(d):
    """SVG path → list of (cmd, [abs points]) with cmd in M L C Q Z (absolute). Handles m l c q h v z, any case."""
    toks = re.findall(r"[MmLlHhVvCcSsQqTtAaZz]|-?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?", d)
    out, i, x, y, sx, sy, cmd = [], 0, 0.0, 0.0, 0.0, 0.0, None
    def num():
        nonlocal i
        v = float(toks[i]); i += 1; return v
    while i < len(toks):
        t = toks[i]
        if t.isalpha():
            cmd = t; i += 1
            if cmd in "Zz":
                out.append(("Z", [])); x, y = sx, sy; continue
        if cmd is None:
            i += 1; continue
        rel = cmd.islower(); c = cmd.upper()
        if c == "M":
            a, b = num(), num(); x, y = (x + a, y + b) if rel else (a, b); sx, sy = x, y
            out.append(("M", [(x, y)])); cmd = "l" if rel else "L"
        elif c == "L":
            a, b = num(), num(); x, y = (x + a, y + b) if rel else (a, b); out.append(("L", [(x, y)]))
        elif c == "H":
            a = num(); x = x + a if rel else a; out.append(("L", [(x, y)]))
        elif c == "V":
            a = num(); y = y + a if rel else a; out.append(("L", [(x, y)]))
        elif c == "C":
            v = [num() for _ in range(6)]
            pts = [(x + v[0], y + v[1]), (x + v[2], y + v[3]), (x + v[4], y + v[5])] if rel else [(v[0], v[1]), (v[2], v[3]), (v[4], v[5])]
            out.append(("C", pts)); x, y = pts[-1]
        elif c == "Q":
            v = [num() for _ in range(4)]
            pts = [(x + v[0], y + v[1]), (x + v[2], y + v[3])] if rel else [(v[0], v[1]), (v[2], v[3])]
            out.append(("Q", pts)); x, y = pts[-1]
        else:   # S, T, A: rare in Slides output; skip their numbers
            n = {"S": 4, "T": 2, "A": 7}[c]; v = [num() for _ in range(n)]
            x, y = (x + v[-2], y + v[-1]) if rel else (v[-2], v[-1]); out.append(("L", [(x, y)]))
    return out


def bbox(segs):
    xs = [p[0] for _, ps in segs for p in ps]; ys = [p[1] for _, ps in segs for p in ps]
    return (min(xs), min(ys), max(xs), max(ys)) if xs else (0, 0, 0, 0)


def fmt(v):
    return f"{v:.4f}".rstrip("0").rstrip(".")


def unit_path(segs, b):
    """Path in 0..1 units of bbox b (data-path)."""
    x0, y0, x1, y1 = b; w, h = (x1 - x0) or 1, (y1 - y0) or 1
    out = []
    for c, ps in segs:
        if c == "Z": out.append("Z"); continue
        out.append(c + " ".join(f"{fmt((px - x0) / w)},{fmt((py - y0) / h)}" for px, py in ps))
    return " ".join(out)


def abs_path(segs):
    return " ".join("Z" if c == "Z" else c + " ".join(f"{fmt(x)},{fmt(y)}" for x, y in ps) for c, ps in segs)


def classify(segs):
    """('rect'|'roundrect'|'ellipse'|'line'|None, extra). Axis-aligned shapes only; anything else is drawn exactly."""
    body = [c for c, _ in segs if c != "M"]
    pts = [p for _, ps in segs for p in ps]
    x0, y0, x1, y1 = bbox(segs); w, h = x1 - x0, y1 - y0
    tol = 0.05 + 0.002 * max(w, h)
    near = lambda a, b: abs(a - b) <= tol
    on_box = all((near(x, x0) or near(x, x1)) and (near(y, y0) or near(y, y1)) for x, y in pts)
    if len(segs) == 2 and segs[0][0] == "M" and segs[1][0] == "L":
        return "line", (segs[0][1][0], segs[1][1][0])
    if set(body) <= {"L", "Z"} and on_box and 4 <= sum(1 for c in body if c == "L") <= 5 and w > 0 and h > 0:
        return "rect", None
    if body.count("C") == 4 and set(body) <= {"C", "Z"} and w > 0 and h > 0:
        ends = [ps[-1] for c, ps in segs if c in "MC"]
        if all(near(x, x0) or near(x, x1) or near(y, y0) or near(y, y1) for x, y in ends):
            return "ellipse", None
    if body.count("C") == 4 and set(body) <= {"C", "L", "Z"} and w > 0 and h > 0:
        rs = []
        prev = segs[0][1][0]
        for c, ps in segs[1:]:
            if c == "C": rs.append(max(abs(ps[-1][0] - prev[0]), abs(ps[-1][1] - prev[1])))
            if ps: prev = ps[-1]
        if rs and max(rs) - min(rs) <= tol * 4 and all(near(x, x0) or near(x, x1) or near(y, y0) or near(y, y1) for x, y in pts if True):
            r = sum(rs) / len(rs)
            return "roundrect", min(0.5, r / max(min(w, h), 1e-6))
    return None, None


# ---------------------------------------------------------------- one slide

def mat_mul(m, n):
    a, b, c, d, e, f = m; A, B, C, D, E, F = n
    return (a * A + c * B, b * A + d * B, a * C + c * D, b * C + d * D, a * E + c * F + e, b * E + d * F + f)


def parse_transform(t):
    m = (1, 0, 0, 1, 0, 0)
    for name, args in re.findall(r"(\w+)\(([^)]*)\)", t or ""):
        v = [float(x) for x in re.split(r"[ ,]+", args.strip()) if x]
        if name == "matrix": n = tuple(v)
        elif name == "translate": n = (1, 0, 0, 1, v[0], v[1] if len(v) > 1 else 0)
        elif name == "scale": n = (v[0], 0, 0, v[1] if len(v) > 1 else v[0], 0, 0)
        else: continue
        m = mat_mul(m, n)
    return m


def apply(m, x, y):
    a, b, c, d, e, f = m
    return a * x + c * y + e, b * x + d * y + f


def extract(svg_text):
    """Walk one slide's SVG in paint order → (viewBox w, h, background, objects)."""
    root = ET.fromstring(svg_text)
    vb = [float(v) for v in root.get("viewBox", "0 0 960 540").split()]
    W, H = vb[2], vb[3]
    clips = {}
    for cp in root.iter(SVG_NS + "clipPath"):
        p = cp.find(SVG_NS + "path")
        if p is not None: clips[cp.get("id")] = bbox(parse_path(p.get("d", "")))
    objs, bg = [], None
    state: dict = {"frame": None, "line": None, "shape": None}

    def full_slide(b):
        return b[0] <= 0.5 and b[1] <= 0.5 and b[2] >= W - 0.5 and b[3] >= H - 0.5

    def visit(el, m, in_label=None):
        nonlocal bg
        tag = el.tag.replace(SVG_NS, "")
        if tag == "clipPath": return
        if tag == "g":
            m2 = mat_mul(m, parse_transform(el.get("transform")))
            label = el.get("aria-label")
            if label is not None:
                hit = el.find(SVG_NS + "path")
                hb = bbox(parse_path(hit.get("d", ""))) if hit is not None else None
                fr = {"kind": "text", "label": label, "hit": hb, "runs": [], "shape": None}
                objs.append(fr); state["frame"] = fr; state["line"] = None
                return   # the label group only holds the transparent hit area
            for c in el: visit(c, m2)
            return
        if tag == "image":
            x, y = float(el.get("x", 0)), float(el.get("y", 0)); w, h = float(el.get("width", 0)), float(el.get("height", 0))
            cid = re.search(r"url\(#([^)]+)\)", el.get("clip-path", "") or "")
            cb = clips.get(cid.group(1)) if cid else None
            vx0, vy0, vx1, vy1 = (max(x, cb[0]), max(y, cb[1]), min(x + w, cb[2]), min(y + h, cb[3])) if cb else (x, y, x + w, y + h)
            (X0, Y0), (X1, Y1) = apply(m, vx0, vy0), apply(m, vx1, vy1)
            crop = [(vy0 - y) / h, (x + w - vx1) / w, (y + h - vy1) / h, (vx0 - x) / w] if w and h else [0, 0, 0, 0]
            objs.append({"kind": "image", "href": el.get(XLINK) or el.get("href"), "box": (min(X0, X1), min(Y0, Y1), max(X0, X1), max(Y0, Y1)),
                         "crop": crop, "rot": math.degrees(math.atan2(m[1], m[0]))})
            state["line"] = None
            return
        if tag != "path": return
        d = el.get("d", "")
        if not d: return
        segs = parse_path(d)
        if m != (1, 0, 0, 1, 0, 0): segs = [(c, [apply(m, *p) for p in ps]) for c, ps in segs]
        b = bbox(segs)
        fill, stroke = el.get("fill"), el.get("stroke")
        clear = el.get("fill-opacity") == "0.0" or fill in (None, "none", "transparent")
        rule = el.get("fill-rule")
        sw = float(el.get("stroke-width", 1)) if stroke else 0
        dash = el.get("stroke-dasharray")
        if bg is None and not stroke and not clear and full_slide(b):
            bg = fill; return
        if bg is not None and not stroke and full_slide(b): return
        closed = any(c == "Z" for c, _ in segs)
        if stroke and not clear:   # filled and outlined in one path: an arrowhead
            ln = state["line"]
            if ln is not None:
                ln["heads"].append({"segs": segs, "fill": fill, "stroke": stroke, "sw": sw}); return
            objs.append({"kind": "shape", "segs": segs, "fill": fill, "stroke": stroke, "sw": sw, "dash": dash}); return
        if stroke:   # outline only: the outline of the shape just filled, or a line / curve of its own
            sh = state["shape"]
            if sh is not None and sh.get("d") == d and sh["stroke"] is None:
                sh.update(stroke=stroke, sw=sw, dash=dash, cap=el.get("stroke-linecap")); return
            ln = state["line"]
            if ln is not None and ln["stroke"] is None:
                ln.update(stroke=stroke, sw=sw, dash=dash, vis=segs, cap=el.get("stroke-linecap")); return
            objs.append({"kind": "stroke", "segs": segs, "stroke": stroke, "sw": sw, "dash": dash, "cap": el.get("stroke-linecap")})
            state["shape"] = None; return
        if rule == "evenodd" or (clear and rule != "nonzero"):   # a shape's fill (maybe transparent) starts an object
            if not closed and [c for c, _ in segs] == ["M", "L"]:   # the full extent of a straight line (its visible stroke follows)
                ln = {"kind": "line", "segs": segs, "stroke": None, "heads": [], "vis": None}
                objs.append(ln); state["line"] = ln; state["shape"] = None; return
            if not closed:   # the hit area of a curve or elbow connector: its stroke follows as a path of its own
                state["shape"] = None; state["line"] = None; return
            sh = {"kind": "shape", "segs": segs, "d": d, "fill": None if clear else fill, "stroke": None, "sw": 0, "dash": None}
            fr = state["frame"]
            if fr is not None and fr["shape"] is None and fr["hit"] and not fr["runs"] and all(abs(p - q) < 1.5 for p, q in zip(fr["hit"], b)):
                fr["shape"] = sh   # the shape behind a text frame
            else:
                objs.append(sh)
            state["shape"] = sh; state["line"] = None; return
        if not clear:   # an opaque nonzero fill without outline: glyphs of the current text frame
            fr = state["frame"]
            if fr is not None:
                fr["runs"].append({"fill": fill, "b": b, "d": d}); return
            objs.append({"kind": "shape", "segs": segs, "d": d, "fill": fill, "stroke": None, "sw": 0, "dash": None})

    for child in root:
        visit(child, (1, 0, 0, 1, 0, 0))
    return W, H, bg, objs


# ---------------------------------------------------------------- text fitting (headless Chrome)

FIT_JS = r"""
async ({frames, families}) => {
  await document.fonts.ready;
  const cv = document.createElement('canvas'), ctx = cv.getContext('2d');
  const V = [{b: 0, i: 0}, {b: 1, i: 0}, {b: 0, i: 1}, {b: 1, i: 1}];
  const fontOf = (fam, v, fs) => `${v.i ? 'italic ' : ''}${v.b ? 'bold ' : ''}${fs}px "${fam}"`;
  const REF = 100, cache = new Map();
  const meas = (fam, vi, s) => {   // metrics at REF px; everything scales linearly with the size
    const k = fam + '|' + vi + '|' + s; let m = cache.get(k);
    if (!m) { ctx.font = fontOf(fam, V[vi], REF); const t = ctx.measureText(s);
      m = {l: t.actualBoundingBoxLeft, r: t.actualBoundingBoxRight, a: t.actualBoundingBoxAscent, d: t.actualBoundingBoxDescent, w: t.width,
           fa: t.fontBoundingBoxAscent, fd: t.fontBoundingBoxDescent}; cache.set(k, m); }
    return m;
  };
  const median = a => { const s = [...a].sort((x, y) => x - y); return s.length ? s[s.length >> 1] : 0; };

  function rowsOf(runs) {
    const rs = [...runs].sort((p, q) => (p.b[1] + p.b[3]) - (q.b[1] + q.b[3]));
    const rows = [];
    for (const r of rs) {
      const h = r.b[3] - r.b[1];
      const row = rows.find(w => { const ov = Math.min(w.y1, r.b[3]) - Math.max(w.y0, r.b[1]); return ov > 0.45 * Math.min(h, w.y1 - w.y0); });
      if (row) { row.runs.push(r); row.y0 = Math.min(row.y0, r.b[1]); row.y1 = Math.max(row.y1, r.b[3]); }
      else rows.push({runs: [r], y0: r.b[1], y1: r.b[3]});
    }
    rows.sort((p, q) => p.y0 - q.y0); rows.forEach(w => w.runs.sort((p, q) => p.b[0] - q.b[0]));
    return rows;
  }

  // Assign the frame's text to its glyph runs in reading order for one family and size; returns cost and the split.
  function assign(T, rows, fam, fs, vset) {
    // Rows in order; inside a row, a small dynamic program splits the text over its glyph runs so the summed width
    // error is least, the row ending where text can wrap. Between rows the best end is kept (greedy).
    let p = 0, err = 0, tot = 0; const out = [];
    const isWS = c => c === ' ' || c === '\n' || c === '\t';
    const runCost = (s, W, Hh) => {
      let best = null;
      for (const vi of vset) {
        const m = meas(fam, vi, s), w = (m.l + m.r) * fs / REF, h = (m.a + m.d) * fs / REF;
        const c = Math.abs(w - W) + 0.5 * Math.abs(h - Hh) + (vi ? 0.015 * W + 0.5 : 0);   // bold only when the widths need it
        if (!best || c < best.c) best = {c, vi, w};
      }
      return best;
    };
    for (let ri = 0; ri < rows.length; ri++) {
      let para = false;
      while (p < T.length && isWS(T[p])) { if (T[p] === '\n') para = true; p++; }
      const gl = rows[ri].runs.filter(r => r.kind === 'glyph');
      if (!gl.length) continue;
      const stop = T.indexOf('\n', p) < 0 ? T.length : T.indexOf('\n', p);
      // layer j: map end position → {c, prev, vi}
      let layer = new Map([[p, {c: 0, prev: null, s: null, vi: 0}]]);
      const layers = [];
      for (let j = 0; j < gl.length; j++) {
        const r = gl[j], W = r.b[2] - r.b[0], Hh = r.b[3] - r.b[1], last = j === gl.length - 1;
        const next = new Map();
        for (const [st, cell] of layer) {
          for (let e = st + 1; e <= stop; e++) {
            const s = T.slice(st, e); if (!s.trim()) continue;
            const m0 = meas(fam, 0, s), w0 = (m0.l + m0.r) * fs / REF;
            if (w0 > W * 1.6 + fs) break;
            if (w0 < W * 0.45 - fs) continue;
            if (last && !(e === stop || isWS(T[e]) || T[e - 1] === '-' || T[e - 1] === '/')) continue;
            const rc = runCost(s, W, Hh);
            let c = cell.c + rc.c;
            if (j && cell.s !== null) {   // the gap between this run's ink and the previous one's must fit the text (a space, or none)
              const mp = meas(fam, cell.vi, cell.s), mc = meas(fam, rc.vi, s);
              const exp = (mp.w - mp.r - mc.l) * fs / REF, act = r.b[0] - gl[j - 1].b[2];
              c += Math.abs(exp - act);
            }
            const cur = next.get(e);
            if (!cur || c < cur.c) next.set(e, {c, prev: st, vi: rc.vi, s});
          }
        }
        if (!next.size) break;
        layers.push(next); layer = next;
      }
      const W0 = gl.reduce((s2, r) => s2 + r.b[2] - r.b[0], 0); tot += W0;
      if (layers.length < gl.length) { err += W0; continue; }   // the text doesn't fit these runs at this size
      let end = null, bestc = Infinity;
      for (const [e, cell] of layers[layers.length - 1]) if (cell.c < bestc) { bestc = cell.c; end = e; }
      err += bestc;
      const picks = [];
      for (let j = gl.length - 1, e = end; j >= 0; j--) { const cell = layers[j].get(e); picks.unshift({r: gl[j], s: T.slice(cell.prev, e), vi: cell.vi}); e = cell.prev; }
      picks.forEach((o, k) => out.push({...o, row: ri, para: para && k === 0}));
      p = end;
    }
    err += (T.slice(p).replace(/\s/g, '').length) * fs * 0.5;   // text left over
    return {cost: err / Math.max(tot, 1), out};
  }

    const ca = document.createElement('canvas'), cb = document.createElement('canvas');
  ca.getContext('2d', {willReadFrequently: true}); cb.getContext('2d', {willReadFrequently: true});
  function runOverlap(o, fam, fs, vi) {   // glyph outline vs our text, over one run's box: [intersection, union]
    const [x0, y0, x1, y1] = [o.r.b[0] - 2, o.r.b[1] - 2, o.r.b[2] + 2, o.r.b[3] + 2];
    const sc = Math.min(4, 800 / (x1 - x0), 400 / (y1 - y0));
    const W = Math.max(1, Math.ceil((x1 - x0) * sc)), H = Math.max(1, Math.ceil((y1 - y0) * sc));
    for (const cv2 of [ca, cb]) { cv2.width = W; cv2.height = H; }
    const A = ca.getContext('2d'), B = cb.getContext('2d');
    for (const g of [A, B]) { g.setTransform(sc, 0, 0, sc, -x0 * sc, -y0 * sc); g.fillStyle = '#000'; }
    A.fill(new Path2D(o.r.d));
    const m = meas(fam, vi, o.s), S = v => v * fs / REF;
    B.font = fontOf(fam, V[vi], fs);
    const tx = o.r.b[0] + S(m.l), ty = (o.r.b[3] - S(m.d) + o.r.b[1] + S(m.a)) / 2;
    B.fillText(o.s, tx, ty);
    // measureText's ink box is approximate (italics especially): align the two inks by their pixel extents
    const ext = d => { let l = W, r = -1, t = H, btm = -1;
      for (let y = 0; y < H; y++) for (let x = 0; x < W; x++) if (d[(y * W + x) * 4 + 3] > 96) { if (x < l) l = x; if (x > r) r = x; if (y < t) t = y; if (y > btm) btm = y; }
      return r < 0 ? null : [l, t, r, btm]; };
    const a = A.getImageData(0, 0, W, H).data;
    let b = B.getImageData(0, 0, W, H).data;
    const ea = ext(a), eb = ext(b);
    let dx = 0, dy = 0;
    if (ea && eb) {
      dx = ((ea[0] + ea[2]) - (eb[0] + eb[2])) / 2 / sc; dy = (ea[3] - eb[3]) / sc;   // centre horizontally, baseline-ish (bottom) vertically
      if (Math.abs(dx) > 0.01 || Math.abs(dy) > 0.01) { B.setTransform(1, 0, 0, 1, 0, 0); B.clearRect(0, 0, W, H); B.setTransform(sc, 0, 0, sc, -x0 * sc, -y0 * sc); B.fillText(o.s, tx + dx, ty + dy); b = B.getImageData(0, 0, W, H).data; }
    }
    let inter = 0, uni = 0;
    for (let i = 3; i < a.length; i += 4) { const p = a[i] > 96, q = b[i] > 96; if (p && q) inter++; if (p || q) uni++; }
    return [inter, uni, dx, dy];
  }
  function overlap(glyphs, c) {   // pick each run's weight / slant by shape (each at the size its widths imply), return the frame's overlap
    let I = 0, U = 0; const sizes = [];
    for (const o of c.out) {
      let bestv = null;
      for (let vi = 0; vi < 4; vi++) {
        const m = meas(c.fam, vi, o.s), w = (m.l + m.r) / REF;
        const f = w > 0 ? (o.r.b[2] - o.r.b[0]) / w : c.fs;
        if (!(f > c.fs * 0.7 && f < c.fs * 1.4)) continue;
        const [i, u, dx, dy] = runOverlap(o, c.fam, f, vi);
        if (!bestv || i / Math.max(u, 1) > bestv.r) bestv = {vi, i, u, dx, dy, f, r: i / Math.max(u, 1)};
      }
      if (!bestv) continue;
      o.vi = bestv.vi; I += bestv.i; U += bestv.u; sizes.push([bestv.f, o.r.b[2] - o.r.b[0], o]);
    }
    // one size for the frame: the width-weighted median of the runs' sizes; realign the runs at it
    if (sizes.length) {
      sizes.sort((p, q) => p[0] - q[0]); const tot = sizes.reduce((s2, x) => s2 + x[1], 0); let acc = 0;
      for (const [f, w] of sizes) { acc += w; if (acc >= tot / 2) { c.fs = f; break; } }
      for (const [, , o] of sizes) { const [, , dx, dy] = runOverlap(o, c.fam, c.fs, o.vi); o.dx = dx; o.dy = dy; }
    }
    return U ? I / U : 0;
  }

  const results = [];
  for (const fr of frames) {
    const T = fr.label.replace(/\r/g, '').replace(/\n+$/, '');
    const runs = fr.runs.map(r => ({...r}));
    if (!runs.length || !T.trim()) { results.push(null); continue; }
    const hs = runs.map(r => r.b[3] - r.b[1]);
    const big = median(hs);
    for (const r of runs) {   // thin wide glyph paths are underlines; tiny ones at a row start are bullets
      const w = r.b[2] - r.b[0], h = r.b[3] - r.b[1];
      r.kind = (h < 0.16 * big && w > 3 * h && w > big) ? 'underline' : 'glyph';
    }
    const rows = rowsOf(runs.filter(r => r.kind === 'glyph'));
    for (const row of rows) {   // a bullet: a small first glyph with a gap after it, when the label has no such character there
      const [a, b] = row.runs;
      if (b && (a.b[2] - a.b[0]) < 0.6 * big && (a.b[3] - a.b[1]) < 0.6 * big && b.b[0] - a.b[2] > 0.3 * big) { a.kind = 'bullet'; }
    }
    const fs0 = big / 0.95;
    const cands = [];
    for (const fam of families) {   // per family: the size whose text widths match the glyph runs best
      let b = null;
      for (let f = fs0 * 0.55; f <= fs0 * 1.6; f += Math.max(0.05, fs0 * 0.01)) {
        const a = assign(T, rows, fam, f, [0, 1]);
        if (!b || a.cost < b.cost) b = {...a, fam, fs: f};
      }
      for (let f = b.fs - fs0 * 0.012; f <= b.fs + fs0 * 0.012; f += Math.max(0.01, fs0 * 0.001)) {
        const a = assign(T, rows, fam, f, [0, 1]);   // italics are told apart by shape (overlap), not by width
        if (a.cost < b.cost) b = {...a, fam, fs: f};
      }
      cands.push(b);
    }
    cands.sort((p, q) => p.cost - q.cost);
    // widths alone can't tell a serif from a sans: compare shapes. Rasterize the glyph outlines and each candidate's
    // text drawn where it would sit, and keep the best overlap among the candidates that fit the widths.
    let shortlist = cands.filter(c => c.cost < 0.25);   // every family that fits the widths at all
    if (!shortlist.length) shortlist = cands.slice(0, 3);
    for (const c of shortlist) c.iou = overlap(runs.filter(r => r.kind === 'glyph'), c);
    if (fr.debug) console.log(JSON.stringify(shortlist.map(c => [c.fam, +c.cost.toFixed(3), +c.iou.toFixed(3), +c.fs.toFixed(2), c.out.map(o => o.vi + ':' + o.s).join('|')])));
    const best = shortlist.reduce((p, q) => (q.iou > p.iou + 0.005 || (Math.abs(q.iou - p.iou) <= 0.005 && q.cost < p.cost)) ? q : p);
    const {fam, fs} = best;
    const S0 = v => v * fs / REF;
    // a frame can mix sizes (a title over a subtitle): give a row its own size when its widths say so
    const rowScale = {};
    for (let ri = 0; ri < rows.length; ri++) {
      const items = best.out.filter(o => o.row === ri);
      const r = items.map(o => { const m = meas(fam, o.vi, o.s); return (o.r.b[2] - o.r.b[0]) / Math.max(S0(m.l + m.r), 0.01); }).filter(v => isFinite(v));
      const k = median(r);
      if (items.length && Math.abs(k - 1) > 0.06) {
        rowScale[ri] = k;
        for (const o of items) { const [i2, u2, dx, dy] = runOverlap(o, fam, fs * k, o.vi); o.dx = dx; o.dy = dy; }
      }
    }
    const S = v => v * fs / REF;
    // Lay out: per row its baseline, left origin, advance width; per run its text and variant.
    const R = rows.map((row, ri) => ({ri, items: best.out.filter(o => o.row === ri), bullet: row.runs.find(r => r.kind === 'bullet'), y0: row.y0, y1: row.y1}));
    for (const row of R) {
      if (!row.items.length) continue;
      row.k = rowScale[row.ri] || 1;
      const Sr = v => v * fs * row.k / REF;
      const bl = row.items.map(o => { const m = meas(fam, o.vi, o.s); return (o.r.b[3] - Sr(m.d) + o.r.b[1] + Sr(m.a)) / 2 + (o.dy || 0); });
      row.base = median(bl);
      const f0 = row.items[0], m0 = meas(fam, f0.vi, f0.s);
      row.x = f0.r.b[0] + Sr(m0.l) + (f0.dx || 0);
      let x = row.x;   // the runs carry every character of the label, spaces included
      for (const o of row.items) x += Sr(meas(fam, o.vi, o.s).w);
      row.w = x - row.x;
      row.para = row.items[0].para || row.ri === 0;
      if (row.bullet) { const bw = row.bullet.b[2] - row.bullet.b[0], bh = row.bullet.b[3] - row.bullet.b[1];
        row.bchar = bh < 0.45 * bw ? '-' : (/[qQcC]/.test(row.bullet.d) ? '●' : '■');
        const mb = meas(fam, 0, row.bchar); row.bx = row.bullet.b[0] + S(mb.l); }
    }
    const lines = R.filter(r => r.items.length);
    // underline: mark the runs above each underline glyph
    for (const u of runs.filter(r => r.kind === 'underline'))
      for (const row of lines) for (const o of row.items)
        if (o.r.b[0] < u.b[2] && o.r.b[2] > u.b[0] && u.b[1] >= row.base - 0.1 * fs && u.b[1] <= row.base + 0.4 * fs) o.u = true;
    // paragraphs
    const paras = [];
    for (const row of lines) { if (row.para || !paras.length) paras.push([row]); else paras[paras.length - 1].push(row); }
    const pitches = paras.flatMap(p => p.slice(1).map((r, i) => r.base - p[i].base));
    const mf = meas(fam, 0, 'Hg');
    const natural = S(mf.fa + mf.fd);
    const pitch = pitches.length ? median(pitches) : natural;
    const L = pitch / fs;
    const half = (pitch - natural) / 2;
    const top = lines[0].base - half - S(mf.fa);
    // horizontal frame: the text frame's own box, else the extent of the text
    const [fx0, , fx1] = fr.frame || [Math.min(...lines.map(r => Math.min(r.x, r.bx ?? r.x))), 0, Math.max(...lines.map(r => r.x + r.w))];
    const mid = (fx0 + fx1) / 2;
    const out = [];
    let prevBase = null;
    for (const p of paras) {
      const lefts = p.map(r => r.x), rights = p.map(r => r.x + r.w), cents = p.map(r => r.x + r.w / 2);
      const spread = a => Math.max(...a) - Math.min(...a);
      let align = 'left';
      if (p.length > 1) align = spread(lefts) < 1 ? 'left' : spread(cents) < 1 ? 'center' : spread(rights) < 1 ? 'right' : 'left';
      else { const lg = lefts[0] - fx0, rg = fx1 - rights[0]; align = (Math.abs(lg - rg) < Math.max(2, 0.08 * fs) && lg > 2) ? 'center' : (rg < lg - 2 && rg < 0.6 * fs) ? 'right' : 'left'; }
      let ml = 0, mr = 0;
      const indent = p[0].bullet ? (Math.min(...lefts) - fx0) : 0;
      if (align === 'left') ml = (p[0].bullet ? p[0].bx : Math.min(...lefts)) - fx0;
      else if (align === 'right') mr = fx1 - Math.max(...rights);
      else { const c = cents.reduce((s, v) => s + v, 0) / cents.length; if (c > mid) ml = 2 * (c - mid); else mr = 2 * (mid - c); }
      const gap = prevBase === null ? 0 : (p[0].base - prevBase) - pitch;
      prevBase = p[p.length - 1].base;
      const rowsHtml = p.map((r, k) => {
        const runsHtml = r.items.map(o => {
          const st = []; const v = V[o.vi];
          if (v.b) st.push('font-weight: bold'); if (v.i) st.push('font-style: italic');
          if (o.r.fill && o.r.fill.toLowerCase() !== fr.color) st.push(`color: ${o.r.fill}`);
          if (o.u) st.push('text-decoration: underline');
          const t = o.s.replace(/&/g, '&amp;').replace(/</g, '&lt;').replace(/>/g, '&gt;').replace(/  /g, ' &nbsp;');
          return st.length ? `<span style="${st.join('; ')}">${t}</span>` : t;
        }).join('');
        return r.k !== 1 ? `<span style="font-size: ${fs * r.k}PT">${runsHtml}</span>` : runsHtml;
      });
      let inner = rowsHtml.join('<br>');
      const st = [];
      if (align !== 'left') st.push(`text-align: ${align}`);
      if (p[0].bullet) {
        const bulletIndent = Math.min(...lefts) - p[0].bx;
        st.push(`padding-left: ${(Math.min(...lefts) - fx0)}PT`);
        inner = `<span style="display: inline-block; width: ${bulletIndent}PT; margin-left: -${bulletIndent}PT">${p[0].bchar}</span>` + inner;
      } else {
        if (ml > 0.05) st.push(`margin-left: ${ml}PT`);
        if (mr > 0.05) st.push(`margin-right: ${mr}PT`);
      }
      if (Math.abs(gap) > 0.3) st.push(`margin-top: ${gap}PT`);
      out.push(`<div${st.length ? ` style="${st.join('; ')}"` : ''}>${inner}</div>`);
    }
    const lastBase = lines[lines.length - 1].base;
    const bottom = lastBase - half - S(mf.fa) + pitch;
    results.push({fam, fs, L, top, bottom, x0: fx0, x1: fx1, html: out.join(''), cost: best.cost, iou: best.iou});
  }
  return results;
}
"""


def fit_text(frames, families):
    """Run FIT_JS over the frames in headless Chrome with every candidate family loaded."""
    from playwright.sync_api import sync_playwright
    google = [f for f in families if f in GOOGLE]
    q = "&".join("family=" + f.replace(" ", "+") + ":ital,wght@0,400;0,700;1,400;1,700" for f in google)
    head = f'<link rel="stylesheet" href="https://fonts.googleapis.com/css2?{q}&display=block">' if google else ""
    loads = [f'document.fonts.load("{st} {wt} 40px \\"{f}\\"")' for f in families for st in ("normal", "italic") for wt in ("400", "700")]
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.set_content(f"<!doctype html><html><head>{head}</head><body></body></html>", wait_until="networkidle")
        pg.evaluate(f"Promise.all([{','.join(loads)}]).catch(() => 0)")
        missing = pg.evaluate("""fams => fams.filter(f => { const c = document.createElement('canvas').getContext('2d');
            const w = font => { c.font = font; return c.measureText('Hamburgefonstiv 0123').width; };
            return w(`40px "${f}", monospace`) === w('40px monospace') && w(`40px "${f}", serif`) === w('40px serif'); })""", families)
        if missing:
            print("fonts not available (skipped):", ", ".join(missing), file=sys.stderr)
            families = [f for f in families if f not in missing]
        pg.on("console", lambda m: print(m.text, file=sys.stderr))
        res = pg.evaluate(FIT_JS, {"frames": frames, "families": families})
        b.close()
    return res


# ---------------------------------------------------------------- preset matching

# A filled outline that isn't a rect / roundrect / ellipse may still be one of the editor's presets (a block arrow, callout,
# star …) drawn at some size and adj values. Rasterise it and every preset (shapes.js, so the geometry is the editor's own)
# into its bounding box, keep the best overlap, refine its adj values one at a time, and accept it above MATCH_IOU; otherwise
# the shape stays an exact SVG box.
MATCH_IOU = 0.96
LOW_FIT = 0.9   # text fits under this glyph overlap go to DECK/source/convert-log.txt

_sj = (ROOT / "editor" / "deck" / "shapes.js").read_text()
_po = _sj[_sj.index("PRESET_OF = {"):_sj.index("};", _sj.index("PRESET_OF = {"))]
PRESET_NAMES = [n for n in dict.fromkeys(["rect", "roundrect", "ellipse"] + re.findall(r"['\"]?([\w-]+)['\"]?:\s*'\w+'", _po))   # simplest first
                if n not in ("line", "arrowline", "elbow", "curve", "path", "wordart", "icon")]

MATCH_JS = r"""({ targets, names }) => {
  const N = 96, cv = new OffscreenCanvas(N, N), x = cv.getContext('2d', { willReadFrequently: true });
  const mask = (draw, flip = '') => { x.setTransform(1, 0, 0, 1, 0, 0); x.clearRect(0, 0, N, N);
    x.setTransform(flip.includes('h') ? -1 : 1, 0, 0, flip.includes('v') ? -1 : 1, flip.includes('h') ? N : 0, flip.includes('v') ? N : 0); draw();
    const d = x.getImageData(0, 0, N, N).data, m = new Uint8Array(N * N); for (let i = 0; i < N * N; i++) m[i] = d[i * 4 + 3] > 127; return m; };
  const iou = (a, b) => { let i = 0, u = 0; for (let k = 0; k < a.length; k++) { i += a[k] & b[k]; u += a[k] | b[k]; } return u ? i / u : 0; };
  const box = document.createElement('div'); box.className = 'box'; document.body.append(box);
  const shape = (name, w, h, adj, flip) => {
    box.dataset.shape = name; for (const k of Object.keys(box.dataset)) if (k.startsWith('adj')) delete box.dataset[k];
    adj.forEach((v, i) => { if (v != null) box.dataset[i ? 'adj' + (i + 1) : 'adj'] = v; });
    box.style.width = w + 'px'; box.style.height = h + 'px'; renderShape(box);
    const ps = [...box.querySelectorAll('svg.shape path.f')].filter(p => p.style.fill !== 'none');
    return mask(() => { x.scale(N / w, N / h); ps.forEach(p => x.fill(new Path2D(p.getAttribute('d')))); }, flip);
  };
  return targets.map(t => {
    const [x0, y0, x1, y1] = t.bbox, w = x1 - x0, h = y1 - y0;
    if (w < 2 || h < 2) return null;
    const want = mask(() => { x.scale(N / w, N / h); x.translate(-x0, -y0); x.fill(new Path2D(t.d), 'evenodd'); });
    const tries = [];
    for (const name of names) for (const flip of ['', 'h', 'v', 'hv']) tries.push({ name, flip, adj: [], s: iou(want, shape(name, w, h, [], flip)) });
    tries.sort((a, b) => b.s - a.s);
    const refine = [...new Set(tries.slice(0, 4).concat(tries.filter(c => names.indexOf(c.name) < 3 && !c.flip)))];
    for (const c of refine) {   // refine adj1..adj4 of the closest few and of rect / roundrect / ellipse, coarse then fine
      const adj = [null, null, null, null]; let s = c.s;
      for (const step of [0.1, 0.02]) for (let i = 0; i < 4; i++) {
        const base = adj[i] ?? null;
        for (let k = -8; k <= 8; k++) {
          const v = base == null ? (step === 0.1 ? (k + 8) / 16 * 1.6 - 0.3 : null) : +(base + k * step).toFixed(3);
          if (v == null) continue;
          const a = adj.slice(); a[i] = v;
          const si = iou(want, shape(c.name, w, h, a, c.flip));
          if (si > s + 1e-4) { s = si; adj[i] = v; }
        }
      }
      c.s = s; c.adj = adj;
    }
    // a simpler shape wins a near tie (a callout with its tail tucked in is just a rectangle)
    const rank = c => c.s + (names.indexOf(c.name) < 3 ? 0.01 : 0) + (c.flip ? 0 : 0.002);
    const pick = refine.sort((a, b) => rank(b) - rank(a))[0];

    return pick && { name: pick.name, flip: pick.flip, adj: pick.adj || [], s: pick.s };

  });
}"""


def match_presets(targets):
    """[{d, bbox}] in slide units → [{name, flip, adj, s} | None]: the editor preset each filled outline is, if any."""
    if not targets: return []
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        pg.set_content("<!doctype html><html><body></body></html>")
        pg.add_script_tag(content=_sj)
        res = pg.evaluate(MATCH_JS, {"targets": targets, "names": PRESET_NAMES})
        b.close()
    return [r if r and r["s"] >= MATCH_IOU else None for r in res]


# ---------------------------------------------------------------- the .pptx: notes and builds

def pptx_extras(pptx: Path, W0: float):
    """Per slide (1-based): {"notes": str, "builds": [(bbox in capture units, click index)]}. The embed SVG carries neither;
    Slides' own .pptx download (File → Download → PowerPoint) has both. Builds are the main sequence's entrance effects:
    each click starts a new step; "with" / "after" effects join the step before."""
    from pptx import Presentation
    prs = Presentation(str(pptx)); k = W0 / prs.slide_width
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    out = {}
    for n, s in enumerate(prs.slides, 1):
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        boxes = {}
        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6: walk(sh.shapes)   # a group: its members can animate on their own
                if sh.left is not None: boxes[str(sh.shape_id)] = (sh.left * k, sh.top * k, (sh.left + sh.width) * k, (sh.top + sh.height) * k)
        walk(s.shapes)
        builds, step = [], -1
        for ctn in s.element.iter(P + "cTn"):
            kind, cls = ctn.get("nodeType"), ctn.get("presetClass")
            if kind not in ("clickEffect", "withEffect", "afterEffect") or cls != "entr": continue
            if kind == "clickEffect" or step < 0: step += 1
            for tgt in ctn.iter(P + "spTgt"):
                if tgt.get("spid") in boxes: builds.append((boxes[tgt.get("spid")], step))
        out[n] = {"notes": notes, "builds": builds}
    return out


def overlap(a, b):
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0])); iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    u = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - ix * iy
    return ix * iy / u if u > 0 else 0


def apply_builds(lines, builds, k):
    """Mark the boxes each build animates as reveal fragments (data-fragment-index = its click), matching by geometry."""
    geo = []
    for i, ln in enumerate(lines):
        m = re.search(r'style="left: (-?[\d.]+)px; top: (-?[\d.]+)px; width: ([\d.]+)px; height: ([\d.]+)px;', ln)
        if m and ln.lstrip().startswith('<div class="box'):
            x, y, w_, h_ = (float(v) / k for v in m.groups()); geo.append((i, (x, y, x + w_, y + h_)))
    for b, step in builds:
        hits = [i for i, g in geo if overlap(g, b) > 0.6 or (g[0] >= b[0] - 1 and g[1] >= b[1] - 1 and g[2] <= b[2] + 1 and g[3] <= b[3] + 1)]
        for i in hits:   # everything inside a built group or text box builds with it
            if 'class="box fragment' in lines[i]: continue
            lines[i] = lines[i].replace('class="box', 'class="box fragment fade-in', 1).replace('<div class="box fragment fade-in', f'<div data-fragment-index="{step}" class="box fragment fade-in', 1)
    return lines


# ---------------------------------------------------------------- deck output


def esc(s):
    return html.escape(s, quote=True)


class Writer:
    def __init__(self, deck: Path, k: float):
        self.deck, self.k, self.n = deck, k, 0

    def px(self, v):
        return f"{v * self.k:.1f}".rstrip("0").rstrip(".")

    def box_style(self, b, extra=""):
        x0, y0, x1, y1 = b
        return f"left: {self.px(x0)}px; top: {self.px(y0)}px; width: {self.px(x1 - x0)}px; height: {self.px(y1 - y0)}px;" + (" " + extra if extra else "")

    def nid(self, slide):
        self.n += 1
        return f"g{slide}-{self.n}"


def dash_name(dash, sw):
    if not dash or not sw: return None
    v = [float(x) / sw for x in re.split(r"[ ,]+", dash.strip()) if x]
    pats = {"dot": [1, 1], "dash": [4, 3], "dashdot": [4, 3, 1, 3], "longdash": [8, 3], "longdashdot": [8, 3, 1, 3]}
    def dist(p):
        q = (p * ((len(v) // len(p)) + 1))[:len(v)]
        return sum(abs(a - b) for a, b in zip(v, q)) + 5 * abs(len(v) - len(p) * (len(v) // len(p) or 1))
    return min(pats, key=lambda n: dist(pats[n]))


def svg_box(w, slide, segs, fill, stroke, sw, dash, extra_attrs=""):
    """A box holding the exact outline as inline SVG (viewBox in the slide's own units)."""
    x0, y0, x1, y1 = bbox(segs)
    pad = sw / 2
    b = (x0 - pad, y0 - pad, x1 + pad, y1 + pad)
    vb = f"{fmt(b[0])} {fmt(b[1])} {fmt(b[2] - b[0])} {fmt(b[3] - b[1])}"
    attrs = [f'd="{abs_path(segs)}"', f'fill="{fill or "none"}"', 'fill-rule="evenodd"']
    if stroke: attrs += [f'stroke="{stroke}"', f'stroke-width="{fmt(sw)}"', 'stroke-linejoin="round"']
    if dash: attrs.append(f'stroke-dasharray="{dash}"')
    svg = (f'<svg viewBox="{vb}" preserveAspectRatio="none" style="display: block; width: 100%; height: 100%; overflow: visible">'
           f'<path {" ".join(attrs)}/></svg>')
    return (f'  <div class="box freeform" id="{w.nid(slide)}"{extra_attrs} style="{w.box_style(b)}">'
            f'<div class="content" style="padding: 0; height: 100%">{svg}</div></div>')


def shape_box(w, slide, sh, text=None):
    segs = sh["segs"]; kind, extra = classify(segs)
    fill, stroke, sw = sh.get("fill"), sh.get("stroke"), sh.get("sw") or 0
    dn = dash_name(sh.get("dash"), sw)
    b = bbox(segs)
    pm = sh.get("preset")
    if kind in ("rect", "roundrect", "ellipse") or pm:
        vars_ = [f"--fill: {fill or 'transparent'}", f"--stroke: {stroke or 'transparent'}", f"--sw: {w.px(sw)}px"]
        if pm and pm["name"] == "roundrect": pm = {**pm, "adj": [min(0.5, pm["adj"][0]) if pm["adj"][0] is not None else None]}
        if pm:   # a block arrow / callout / star … as the editor's own preset, so it stays resizable and adjustable

            keys = ["adj", "adj2", "adj3", "adj4"]
            attrs = f' data-shape="{pm["name"]}"' + "".join(f' data-{k}="{v:.3f}"' for k, v in zip(keys, pm["adj"]) if v is not None)
            attrs += f' data-flip="{pm["flip"]}"' if pm["flip"] else ""
        else:
            attrs = f' data-shape="{kind}"' + (f' data-adj="{extra:.4f}"' if kind == "roundrect" else "")
        attrs += f' data-dash="{dn}"' if dn else ""

        content, cstyle = "", ""
        if text:
            content, cstyle = text
        top = ' data-valign="top"' if text else ''   # (no backslashes inside an f-string's braces before Python 3.12)
        return (f'  <div class="box" id="{w.nid(slide)}"{attrs}{top} style="{w.box_style(b, "; ".join(vars_) + ";")}">'
                f'<div class="content"{f" style={chr(34)}{cstyle}{chr(34)}" if cstyle else ""}>{content}</div></div>')
    if not fill:   # an outline without fill: the editor's own path shape, exact
        return path_box(w, slide, segs, stroke, sw, dn, closed=True) + (text_overlay(w, slide, b, text) if text else "")
    return svg_box(w, slide, segs, fill, stroke, sw, sh.get("dash")) + (text_overlay(w, slide, b, text) if text else "")


def text_overlay(w, slide, b, text):
    content, cstyle = text
    return (f'\n  <div class="box" id="{w.nid(slide)}" data-valign="top" style="{w.box_style(b)}">'
            f'<div class="content" style="{cstyle}">{content}</div></div>')


def path_box(w, slide, segs, stroke, sw, dn, closed=False):
    b = bbox(segs)
    if b[2] - b[0] < 0.01: b = (b[0] - 0.005, b[1], b[2] + 0.005, b[3])
    if b[3] - b[1] < 0.01: b = (b[0], b[1] - 0.005, b[2], b[3] + 0.005)
    vars_ = f"--stroke: {stroke or 'transparent'}; --sw: {w.px(sw)}px;"
    return (f'  <div class="box" id="{w.nid(slide)}" data-shape="path" data-path="{unit_path(segs, b)}"'
            f'{f" data-dash={chr(34)}{dn}{chr(34)}" if dn else ""} style="{w.box_style(b, vars_)}"><div class="content"></div></div>')


def line_html(w, slide, ln):
    """A straight line: its visible stroke as a line box, each arrowhead exact, all three grouped."""
    vis = ln.get("vis") or ln["segs"]
    (ax, ay), (bx, by) = vis[0][1][0], vis[-1][1][-1]
    sw = ln.get("sw") or 1
    dn = dash_name(ln.get("dash"), sw)
    x0, y0, x1, y1 = min(ax, bx), min(ay, by), max(ax, bx), max(ay, by)
    up = (bx - ax) * (by - ay) < 0
    vars_ = f"--stroke: {ln.get('stroke') or '#000000'}; --sw: {w.px(sw)}px;"
    attrs = ' data-shape="line"' + (' data-dir="up"' if up else "") + (f' data-dash="{dn}"' if dn else "")
    if not ln["heads"]:
        return f'  <div class="box" id="{w.nid(slide)}"{attrs} style="{w.box_style((x0, y0, x1, y1), vars_)}"><div class="content"></div></div>'
    parts = [((x0, y0, x1, y1), f'<div class="box" id="{w.nid(slide)}"{attrs} style="@@; {vars_}"><div class="content"></div></div>')]
    for h in ln["heads"]:
        hb = bbox(h["segs"]); pad = h["sw"] / 2; hb = (hb[0] - pad, hb[1] - pad, hb[2] + pad, hb[3] + pad)
        vb = f"{fmt(hb[0])} {fmt(hb[1])} {fmt(hb[2] - hb[0])} {fmt(hb[3] - hb[1])}"
        svg = (f'<svg viewBox="{vb}" preserveAspectRatio="none" style="display: block; width: 100%; height: 100%; overflow: visible">'
               f'<path d="{abs_path(h["segs"])}" fill="{h["fill"]}" stroke="{h["stroke"]}" stroke-width="{fmt(h["sw"])}" stroke-linejoin="miter"/></svg>')
        parts.append((hb, f'<div class="box freeform" id="{w.nid(slide)}" style="@@"><div class="content" style="padding: 0; height: 100%">{svg}</div></div>'))
    gx0 = min(p[0][0] for p in parts); gy0 = min(p[0][1] for p in parts); gx1 = max(p[0][2] for p in parts); gy1 = max(p[0][3] for p in parts)
    inner = []
    for b, h in parts:
        rel = (b[0] - gx0, b[1] - gy0, b[2] - gx0, b[3] - gy0)
        inner.append("    " + h.replace("@@", w.box_style(rel).rstrip(";")))
    return f'  <div class="box group" id="{w.nid(slide)}" style="{w.box_style((gx0, gy0, gx1, gy1))}">\n' + "\n".join(inner) + "\n  </div>"


def save_image(deck: Path, href: str, images: dict, slide: int):
    if not href: return None
    data = None
    src = images.get(href) or images.get(href.replace("&amp;", "&"))
    if src and src.startswith("data:"):
        head, b64 = src.split(",", 1); data = base64.b64decode(b64); mime = head[5:].split(";")[0]
    else:
        try:
            req = urllib.request.Request(href.replace("&amp;", "&"), headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=30) as r:
                data = r.read(); mime = r.headers.get_content_type()
        except Exception:
            return None
    ext = {"image/png": ".png", "image/jpeg": ".jpg", "image/gif": ".gif", "image/webp": ".webp", "image/svg+xml": ".svg"}.get(mime, ".png")
    name = f"s{slide:02d}-{hashlib.sha1(data).hexdigest()[:10]}{ext}"
    (deck / "assets").mkdir(parents=True, exist_ok=True)
    (deck / "assets" / name).write_bytes(data)
    return f"assets/{name}"


def convert(capture: Path, deck: Path, only=None, keep=False, pptx: Path | None = None):
    cap = json.loads(capture.read_text())
    images = cap.get("images", {})
    deck.mkdir(parents=True, exist_ok=True)
    nums = only or list(range(1, len(cap["slides"]) + 1))
    slides = {n: extract(cap["slides"][n - 1]) for n in nums}
    W0, H0 = slides[nums[0]][0], slides[nums[0]][1]
    k = 1920 / W0
    # text frames → fit fonts in one browser session
    frames, where = [], []
    for n in nums:
        W, H, bg, objs = slides[n]
        for o in objs:
            if o["kind"] == "text" and o["runs"]:
                hit = o["shape"] and bbox(o["shape"]["segs"]) or o["hit"]
                fills = [r["fill"].lower() for r in o["runs"]]
                o["color"] = max(set(fills), key=fills.count)
                frames.append({"label": html.unescape(o["label"]), "runs": [{"b": r["b"], "fill": r["fill"].lower(), "d": r["d"]} for r in o["runs"]],
                               "frame": list(hit) if hit else None, "color": o["color"],
                               "debug": bool(os.environ.get("GS_DEBUG")) and os.environ["GS_DEBUG"] in html.unescape(o["label"])})
                where.append(o); o["slide"] = n
    families = WEBSAFE + GOOGLE
    fits = fit_text(frames, families) if frames else []
    for o, f in zip(where, fits): o["fit"] = f
    used_google = sorted({o["fit"]["fam"] for o in where if o.get("fit") and o["fit"]["fam"] in GOOGLE})

    # filled outlines that aren't rect / roundrect / ellipse: try the editor's presets
    odd = [sh for n in nums for o in slides[n][3] for sh in [o if o["kind"] == "shape" else o.get("shape") if o["kind"] == "text" else None]
           if sh and sh.get("fill") and classify(sh["segs"])[0] not in ("rect", "roundrect", "ellipse", "line")]
    for sh, m in zip(odd, match_presets([{"d": abs_path(sh["segs"]), "bbox": list(bbox(sh["segs"]))} for sh in odd])):
        if m: sh["preset"] = m
    if odd: print(f"{sum(1 for sh in odd if sh.get('preset'))} of {len(odd)} filled outlines matched an editor preset")

    extras = pptx_extras(pptx, W0) if pptx else {}
    if pptx: print(f"from {pptx.name}: notes on {sum(1 for e in extras.values() if e['notes'])} slides, builds on {sum(1 for e in extras.values() if e['builds'])}")
    w = Writer(deck, k)
    sections = []

    for n in nums:
        W, H, bg, objs = slides[n]
        w.n = 0
        out = []
        on_slide = lambda b: not (b[2] < 0 or b[0] > W or b[3] < 0 or b[1] > H)
        for o in objs:
            if o["kind"] == "text":
                f = o.get("fit")
                sh = o["shape"]
                if sh and not sh.get("fill") and not sh.get("stroke"): sh = None   # a plain text box
                if not f:
                    if sh and on_slide(bbox(sh["segs"])): out.append(shape_box(w, n, sh))
                    continue
                pxify = lambda s: re.sub(r"(-?[\d.]+)PT", lambda m: w.px(float(m.group(1))) + "px", s)
                body = pxify(f["html"])
                cstyle = (f"font-family: '{f['fam']}'{', serif' if f['fam'] in ('Georgia', 'Times New Roman', 'Merriweather', 'Lora', 'Playfair Display', 'Roboto Slab', 'EB Garamond') else ', sans-serif'}; "
                          f"font-size: {w.px(f['fs'])}px; line-height: {f['L']:.4f}; color: {o['color']}; text-align: left")
                if sh:
                    sb = bbox(sh["segs"])
                    if not on_slide(sb): continue
                    pad_top = f["top"] - sb[1]
                    cs = cstyle + f"; margin: 0; padding: {w.px(pad_top)}px 0 0 0; width: 100%"
                    out.append(shape_box(w, n, sh, (body, cs)))
                else:
                    b = (f["x0"], f["top"], f["x1"], f["bottom"])
                    if not on_slide(b): continue
                    out.append(f'  <div class="box" id="{w.nid(n)}" data-valign="top" style="{w.box_style(b)}">'
                               f'<div class="content" style="{cstyle}; padding: 0">{body}</div></div>')
            elif o["kind"] == "shape":
                if on_slide(bbox(o["segs"])): out.append(shape_box(w, n, o))
            elif o["kind"] == "stroke":
                if on_slide(bbox(o["segs"])): out.append(path_box(w, n, o["segs"], o["stroke"], o["sw"], dash_name(o.get("dash"), o["sw"])))
            elif o["kind"] == "line":
                if on_slide(bbox(o["segs"])): out.append(line_html(w, n, o))
            elif o["kind"] == "image":
                b = o["box"]
                if not on_slide(b): continue
                src = save_image(deck, o["href"], images, n)
                crop = o["crop"]
                attrs = f' data-crop="{" ".join(f"{c:.4f}" for c in crop)}"' if any(c > 0.001 for c in crop) else ""
                rot = f" transform: rotate({o['rot']:.2f}deg);" if abs(o["rot"]) > 0.05 else ""
                img = f'<img src="{src}" alt="">' if src else ""
                out.append(f'  <div class="box image" id="{w.nid(n)}"{attrs} style="{w.box_style(b, rot.strip())}">'
                           f'<div class="content">{img}</div></div>' + ("" if src else "  <!-- picture not reachable from the capture -->"))
        ex = extras.get(n, {})
        if ex.get("builds"): out = apply_builds(out, ex["builds"], k)
        note = (f'  <aside class="notes">{esc(ex["notes"])}</aside>' if ex.get("notes") else
                f'  <aside class="notes">From Google Slides "{esc(cap.get("title", ""))}" slide {n} (gslides2deck).</aside>')
        sections.append(f'<section data-background-color="{bg or "#ffffff"}">\n' + "\n".join(out + [note]) + "\n</section>")

    path = deck / "deck.html"
    if not path.exists():
        from glideshow.decks import new_deck
        new_deck(path, cap.get("title", ""))
    text = path.read_text()
    joined = "\n".join(sections)
    if keep:
        text = text.replace("\n</div></div>\n<script", "\n" + joined + "\n</div></div>\n<script", 1)
    else:
        text = re.sub(r'(<div class="slides"[^>]*>\n).*?(\n</div></div>)', lambda m: m.group(1) + joined + m.group(2), text, count=1, flags=re.S)
    if H0 * k > 1080.5 or H0 * k < 1079.5:
        text = re.sub(r'<div class="slides"[^>]*>', f'<div class="slides" data-size="1920x{round(H0 * k)}">', text, count=1)
    if used_google:
        link = ('<link id="deck-fonts" rel="stylesheet" href="https://fonts.googleapis.com/css2?' +
                "&".join("family=" + f.replace(" ", "+") + ":ital,wght@0,400;0,700;1,400;1,700" for f in used_google) + '&display=swap">')
        if 'id="deck-fonts"' in text: text = re.sub(r'<link id="deck-fonts"[^>]*>', link, text, count=1)
        else: text = text.replace("</head>", link + "\n</head>", 1)
    path.write_text(text)
    print(f"wrote {path}: {len(nums)} slides, {len(frames)} text frames fitted"
          + (f", Google Fonts: {', '.join(used_google)}" if used_google else ""))
    worst = sorted(((o["fit"]["iou"], o["fit"]["fam"], o["label"][:40]) for o in where if o.get("fit")))[:5]
    if worst: print("least certain text fits (glyph overlap):", "; ".join(f"{c:.2f} {f} {l!r}" for c, f, l in worst))
    # every fit below LOW_FIT for the user to check by eye (font, size or line breaks may be off)
    low = sorted((o for o in where if o.get("fit") and o["fit"]["iou"] < LOW_FIT), key=lambda o: (o["slide"], o["fit"]["iou"]))
    log = deck / "source" / "convert-log.txt"
    if low:
        log.parent.mkdir(parents=True, exist_ok=True)
        log.write_text(f"gslides2deck {capture.name}: text matched with low confidence (glyph overlap < {LOW_FIT}); check these by eye.\n\n"
                       + "".join(f"slide {o['slide']:>3}  {o['fit']['iou']:.2f}  {o['fit']['fam']:<18} {w.px(o['fit']['fs'])} px  {html.unescape(o['label']).strip()[:80]!r}\n" for o in low))
        print(f"{len(low)} low-confidence text fits listed in {log}")
    elif log.exists():
        log.unlink()


# Animated GIFs stay GIFs in the deck; for comparing, both the references and test/snap.py show their first frame
# (createImageBitmap decodes an animated image's first frame), so a screenshot never catches them mid-animation.
FIRST_FRAME_JS = r"""async () => {
  const still = async src => { const bm = await createImageBitmap(await (await fetch(src)).blob()); const c = document.createElement('canvas');
    c.width = bm.width; c.height = bm.height; c.getContext('2d').drawImage(bm, 0, 0); return c.toDataURL('image/png'); };
  for (const im of document.querySelectorAll('img')) if (/^data:image\/gif|\.gif(\?|$)/i.test(im.src)) { im.src = await still(im.src); await im.decode().catch(() => 0); }
  for (const im of document.querySelectorAll('image')) { const h = im.getAttribute('href') || im.getAttribute('xlink:href') || '';
    if (/^data:image\/gif|\.gif(\?|$)/i.test(h)) { const p = await still(h); im.setAttribute('href', p); im.setAttributeNS('http://www.w3.org/1999/xlink', 'xlink:href', p); } }
}"""


def render_refs(capture: Path, out: Path, only=None):
    """Render each captured slide's SVG to OUT/slide-NN.png at 1920 px wide: exact references, no downloads needed."""
    from playwright.sync_api import sync_playwright
    cap = json.loads(capture.read_text()); images = cap.get("images", {})
    out.mkdir(parents=True, exist_ok=True)
    nums = only or list(range(1, len(cap["slides"]) + 1))
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1920, "height": 1080})
        for n in nums:
            svg = cap["slides"][n - 1]
            svg = re.sub(r'xlink:href="([^"]+)"', lambda m: f'xlink:href="{images.get(html.unescape(m.group(1)), m.group(1))}"', svg)
            vb = [float(v) for v in re.search(r'viewBox="([^"]+)"', svg).group(1).split()]
            h = round(1920 * vb[3] / vb[2])
            pg.set_viewport_size({"width": 1920, "height": h})
            pg.set_content('<!doctype html><html><body style="margin:0">' + svg.replace("<svg ", '<svg width="1920" height="%d" ' % h, 1) + "</body></html>")
            pg.evaluate("Promise.all([...document.querySelectorAll('image')].map(i => i.href.baseVal.startsWith('data:') ? 0 : new Promise(r => setTimeout(r, 1500))))")
            pg.evaluate(FIRST_FRAME_JS)
            pg.wait_for_timeout(100)

            pg.screenshot(path=str(out / f"slide-{n:02d}.png"))
        b.close()
    print(f"{len(nums)} reference renders -> {out}")


def selftest():
    """Checks the parts that don't need a capture: preset matching, .pptx notes / builds, fragments by geometry."""
    import tempfile
    from pptx import Presentation
    from pptx.util import Emu
    from lxml import etree
    fails = []
    arrow = "M0,50 L300,50 L300,0 L400,100 L300,200 L300,150 L0,150 Z"   # OOXML rightArrow, adj 50000 / 50000, 400 × 200
    rr = "M20,0 L380,0 C391,0 400,9 400,20 L400,180 C400,191 391,200 380,200 L20,200 C9,200 0,191 0,180 L0,20 C0,9 9,0 20,0 Z"
    blob = "M0,100 C40,-60 120,40 200,0 C300,60 420,-20 400,120 C330,260 120,140 0,100 Z"
    m = match_presets([{"d": arrow, "bbox": [0, 0, 400, 200]}, {"d": rr, "bbox": [0, 0, 400, 200]}, {"d": blob, "bbox": [0, -20, 420, 200]}])
    if not (m[0] and m[0]["name"] == "arrow"): fails.append(f"arrow → {m[0]}")
    if not (m[1] and m[1]["name"] == "roundrect"): fails.append(f"rounded rectangle → {m[1]}")
    if m[2] is not None: fails.append(f"freeform blob → {m[2]}")
    prs = Presentation(); prs.slide_width, prs.slide_height = Emu(9144000), Emu(5143500)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_textbox(Emu(914400), Emu(914400), Emu(1828800), Emu(914400)); b = s.shapes.add_textbox(Emu(4572000), Emu(914400), Emu(1828800), Emu(914400))
    s.notes_slide.notes_text_frame.text = "Say hello first"
    step = lambda i, sp, kind: (f'<p:par><p:cTn id="{i}" fill="hold"><p:childTnLst><p:par><p:cTn id="{i + 1}" presetID="10" presetClass="entr" nodeType="{kind}">'
                                f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{i + 2}" dur="1"/><p:tgtEl><p:spTgt spid="{sp}"/></p:tgtEl></p:cBhvr></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    timing = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" nodeType="tmRoot"><p:childTnLst><p:seq><p:cTn id="2" nodeType="mainSeq"><p:childTnLst>'
              + step(3, a.shape_id, "clickEffect") + step(6, b.shape_id, "clickEffect") + '</p:childTnLst></p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    s.element.append(etree.fromstring(timing))
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "t.pptx"; prs.save(f); ex = pptx_extras(f, 960)
    e = ex.get(1, {})
    if e.get("notes") != "Say hello first": fails.append(f"notes → {e.get('notes')!r}")
    if [st for _, st in e.get("builds", [])] != [0, 1]: fails.append(f"builds → {e.get('builds')}")
    lines = ['  <div class="box" id="g1-1" style="left: 192px; top: 192px; width: 384px; height: 192px;"><div class="content">A</div></div>',
             '  <div class="box" id="g1-2" style="left: 960px; top: 192px; width: 384px; height: 192px;"><div class="content">B</div></div>']
    out = apply_builds(lines, e.get("builds", []), 2)
    if not ('data-fragment-index="0" class="box fragment fade-in"' in out[0] and 'data-fragment-index="1" class="box fragment fade-in"' in out[1]): fails.append(f"fragments → {out}")
    print("selftest ok" if not fails else "selftest FAILED: " + "; ".join(fails))
    return not fails


def parse_range(s):
    out = []
    for part in s.split(","):
        a, _, b = part.partition("-")
        out += list(range(int(a), int(b or a) + 1))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("snippet", help="print the capture snippet to run in a docs.google.com tab")
    c = sub.add_parser("convert", help="capture JSON → deck")
    c.add_argument("capture", type=Path); c.add_argument("deck", type=Path)
    c.add_argument("--slides", type=parse_range); c.add_argument("--keep", action="store_true", help="append instead of replacing slides")
    c.add_argument("--pptx", type=Path, help="the deck's own .pptx download: speaker notes and on-click builds (the capture has neither)")
    r = sub.add_parser("refs", help="render the captured slides themselves to PNG (exact references for comparing)")
    r.add_argument("capture", type=Path); r.add_argument("out", type=Path); r.add_argument("--slides", type=parse_range)
    sub.add_parser("selftest", help="check preset matching and the .pptx notes / builds on synthetic input")
    a = ap.parse_args()
    if a.cmd == "snippet":
        print(SNIPPET.strip())
    elif a.cmd == "selftest":
        sys.exit(0 if selftest() else 1)
    elif a.cmd == "refs":

        render_refs(a.capture, a.out.expanduser(), a.slides)
    else:
        convert(a.capture, a.deck.expanduser(), a.slides, a.keep, a.pptx and a.pptx.expanduser())



if __name__ == "__main__":
    main()
